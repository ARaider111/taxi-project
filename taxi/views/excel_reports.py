from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


HEADER_FILL = PatternFill(
    fill_type="solid",
    start_color="4472C4",
)

HEADER_FONT = Font(
    bold=True,
    color="FFFFFF",
)

HEADER_ALIGNMENT = Alignment(
    horizontal="center",
    vertical="center",
)


def create_workbook(title, headers):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = title
    worksheet.append(headers)

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT

    return workbook, worksheet


def autofit_columns(worksheet, max_width=50):
    for column in worksheet.columns:
        max_length = max(
            len(str(cell.value or ""))
            for cell in column
        )

        letter = column[0].column_letter
        worksheet.column_dimensions[letter].width = min(
            max_length + 2,
            max_width,
        )


def workbook_response(workbook, filename):
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = (
        f'attachment; filename="{filename}"'
    )
    return response