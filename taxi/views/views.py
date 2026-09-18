import openpyxl
from django.http import HttpResponse
from django.contrib.auth.views import LoginView
from django.shortcuts import redirect, render
from django.contrib.auth.decorators import login_required
from taxi.models import User, Driver, Client, Shift, District, Street, Tariff, Order, AuditLog
from django.contrib.auth import logout
from django.db.models import Q, Sum, Count
from django.db import models
from django.utils import timezone
from taxi.utils import log_action 
from openpyxl import Workbook
from django.http import HttpResponse
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime, timedelta
from django.utils import timezone as tz


class CustomLoginView(LoginView):
    template_name = "login.html"

    def form_valid(self, form):
        user = form.get_user()

        if hasattr(user, "is_archive") and user.is_archive:
            form.add_error(None, "Пользователь находится в архиве и не может войти")
            return self.form_invalid(form)

        login_result = super().form_valid(form)

        log_action(
            user=user,
            action="LOGIN",
            description=f"Вход в систему: {user.login}",
            ip_address=self.request.META.get("REMOTE_ADDR"),
        )
        
        if hasattr(user, "role") and user.role == "admin":
            return redirect("admin_panel")

        if hasattr(user, "role") and user.role == "dispatcher":
            return redirect("dispatcher_dashboard")

        return redirect("index")


@login_required
def admin_panel(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    return render(request, "admin_panel.html")


@login_required
def drivers_list(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    drivers = Driver.objects.all()

    query = request.GET.get("q", "").strip()
    if query:
        drivers = drivers.filter(
            Q(lname__icontains=query) |
            Q(fname__icontains=query) |
            Q(phone__icontains=query) |
            Q(car_model__icontains=query) |
            Q(car_number__icontains=query)
        )

    status = request.GET.get("status", "")
    if status:
        drivers = drivers.filter(status=status)

    blacklist = request.GET.get("blacklist", "")
    if blacklist == "1":
        drivers = drivers.filter(is_blacklist=True)
    elif blacklist == "0":
        drivers = drivers.filter(is_blacklist=False)

    archive = request.GET.get("archive", "")
    if archive == "1":
        drivers = drivers.filter(is_archive=True)
    elif archive == "0":
        drivers = drivers.filter(is_archive=False)

    statuses = Driver.STATUS_CHOICES

    return render(request, "drivers_list.html", {
        "drivers": drivers,
        "statuses": statuses,
        "query": query,
        "current_status": status,
        "current_blacklist": blacklist,
        "current_archive": archive,
    })


@login_required
def clients_list(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    clients = Client.objects.all()

    query = request.GET.get("q", "").strip()
    if query:
        clients = clients.filter(
            Q(lname__icontains=query) |
            Q(fname__icontains=query) |
            Q(phone__icontains=query)
        )

    blacklist = request.GET.get("blacklist", "")
    if blacklist == "1":
        clients = clients.filter(is_blacklist=True)
    elif blacklist == "0":
        clients = clients.filter(is_blacklist=False)

    return render(request, "clients_list.html", {
        "clients": clients,
        "query": query,
        "current_blacklist": blacklist,
    })


@login_required
def users_list(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    users = User.objects.all()

    # Поиск по логину, фамилии, имени, отчеству, телефону
    query = request.GET.get("q", "").strip()
    if query:
        users = users.filter(
            Q(login__icontains=query) |
            Q(lname__icontains=query) |
            Q(fname__icontains=query) |
            Q(patronimyc__icontains=query) |
            Q(phone__icontains=query)
        )

    role = request.GET.get("role", "")
    if role:
        users = users.filter(role=role)

    is_active = request.GET.get("is_active", "")
    if is_active == "1":
        users = users.filter(is_active=True)
    elif is_active == "0":
        users = users.filter(is_active=False)

    archive = request.GET.get("archive", "")
    if archive == "1":
        users = users.filter(is_archive=True)
    elif archive == "0":
        users = users.filter(is_archive=False)

    roles = User.objects.values_list("role", flat=True).distinct()

    return render(request, "users_list.html", {
        "users": users,
        "roles": roles,
        "query": query,
        "current_role": role,
        "current_is_active": is_active,
        "current_archive": archive,
    })


@login_required
def shifts_list(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    shifts = Shift.objects.select_related(
        "driver", "opened_user", "closed_user"
    ).order_by("-start_shift")

    query = request.GET.get("q", "").strip()
    if query:
        shifts = shifts.filter(
            Q(driver__lname__icontains=query) |
            Q(driver__fname__icontains=query) |
            Q(start_shift__icontains=query) |
            Q(end_shift__icontains=query)
        )

    status = request.GET.get("status", "")
    if status == "not_opened":
        shifts = shifts.filter(opened_user__isnull=True)
    elif status == "opened":
        shifts = shifts.filter(
            opened_user__isnull=False,
            closed_user__isnull=True
        )
    elif status == "closed":
        shifts = shifts.filter(closed_user__isnull=False)

    return render(request, "shifts_list.html", {
        "shifts": shifts,
        "query": query,
        "current_status": status,
    })
@login_required
def add_driver(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    if request.method == "POST":
        lname = request.POST.get("lname")
        fname = request.POST.get("fname")
        patronimyc = request.POST.get("patronimyc") or None
        phone = request.POST.get("phone")
        car_model = request.POST.get("car_model")
        car_number = request.POST.get("car_number")

        driver = Driver.objects.create(
            lname=lname,
            fname=fname,
            patronimyc=patronimyc,
            phone=phone,
            car_model=car_model,
            car_number=car_number,
            status="Вне работы",
            is_blacklist=False,
        )

        log_action(
            user=request.user,
            action="CREATE",
            object_type="Driver",
            object_id=driver.driver_id,
            description=f"Создан водитель {driver.lname} {driver.fname}",
            ip_address=request.META.get("REMOTE_ADDR"),
        )
        return redirect("drivers_list")

    return render(request, "add_driver.html")

@login_required
def edit_driver(request, driver_id):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    driver = Driver.objects.get(pk=driver_id)

    if request.method == "POST":
        old_data = {
            "lname": driver.lname,
            "fname": driver.fname,
            "patronimyc": driver.patronimyc,
            "phone": driver.phone,
            "car_model": driver.car_model,
            "car_number": driver.car_number,
            "status": driver.status,
        }

        driver.lname = request.POST.get("lname")
        driver.fname = request.POST.get("fname")
        driver.patronimyc = request.POST.get("patronimyc") or None
        driver.phone = request.POST.get("phone")
        driver.car_model = request.POST.get("car_model")
        driver.car_number = request.POST.get("car_number")
        driver.status = request.POST.get("status", driver.status)
        driver.save()

        log_action(
            user=request.user,
            action="UPDATE",
            object_type="Driver",
            object_id=driver.driver_id,
            description=f"Изменены данные водителя {driver.lname} {driver.fname}",
            extra_data={"old_data": old_data},
            ip_address=request.META.get("REMOTE_ADDR"),
        )
        return redirect("drivers_list")

    return render(request, "edit_driver.html", {"driver": driver})

@login_required
def toggle_driver_archive(request, driver_id):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    driver = Driver.objects.get(pk=driver_id)
    old_archive = driver.is_archive
    driver.is_archive = not driver.is_archive  
    driver.save()

    log_action(
        user=request.user,
        action="TOGGLE_ARCHIVE",
        object_type="Driver",
        object_id=driver.driver_id,
        description=f"Архивирование водителя {driver.lname} {driver.fname} ({'в архив' if driver.is_archive else 'из архива'})",
        extra_data={"old_archive": old_archive, "new_archive": driver.is_archive},
        ip_address=request.META.get("REMOTE_ADDR"),
    )

    return redirect("drivers_list")

@login_required
def toggle_driver_blacklist(request, driver_id):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    driver = Driver.objects.get(pk=driver_id)
    old_blacklist = driver.is_blacklist
    driver.is_blacklist = not driver.is_blacklist
    driver.save()

    log_action(
        user=request.user,
        action="TOGGLE_BLACKLIST",
        object_type="Driver",
        object_id=driver.driver_id,
        description=f"Чёрный список водителя {driver.lname} {driver.fname} ({'добавлен' if driver.is_blacklist else 'убран'})",
        extra_data={"old_blacklist": old_blacklist, "new_blacklist": driver.is_blacklist},
        ip_address=request.META.get("REMOTE_ADDR"),
    )

    return redirect("drivers_list")

@login_required
def add_client(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    if request.method == "POST":
        lname = request.POST.get("lname")
        fname = request.POST.get("fname")
        patronimyc = request.POST.get("patronimyc") or None
        phone = request.POST.get("phone")

        client = Client.objects.create(
            lname=lname,
            fname=fname,
            patronimyc=patronimyc,
            phone=phone,
            is_blacklist=False,  
        )

        log_action(
            user=request.user,
            action="CREATE",
            object_type="Client",
            object_id=client.client_id,
            description=f"Создан клиент {client.lname} {client.fname}",
            ip_address=request.META.get("REMOTE_ADDR"),
        )
        return redirect("clients_list")


    return render(request, "add_client.html")

@login_required
def edit_client(request, client_id):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    client = Client.objects.get(pk=client_id)

    if request.method == "POST":
        old_data = {
            "lname": client.lname,
            "fname": client.fname,
            "patronimyc": client.patronimyc,
            "phone": client.phone,
        }

        client.lname = request.POST.get("lname")
        client.fname = request.POST.get("fname")
        client.patronimyc = request.POST.get("patronimyc") or None
        client.phone = request.POST.get("phone")
        client.save()

        log_action(
            user=request.user,
            action="UPDATE",
            object_type="Client",
            object_id=client.client_id,
            description=f"Изменены данные клиента {client.lname} {client.fname}",
            extra_data={"old_data": old_data},
            ip_address=request.META.get("REMOTE_ADDR"),
        )
        return redirect("clients_list")

    return render(request, "edit_client.html", {"client": client})


@login_required
def toggle_client_blacklist(request, client_id):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    client = Client.objects.get(pk=client_id)
    old_blacklist = client.is_blacklist
    client.is_blacklist = not client.is_blacklist
    client.save()

    log_action(
        user=request.user,
        action="TOGGLE_BLACKLIST",
        object_type="Client",
        object_id=client.client_id,
        description=f"Чёрный список клиента {client.lname} {client.fname} ({'добавлен' if client.is_blacklist else 'убран'})",
        extra_data={"old_blacklist": old_blacklist, "new_blacklist": client.is_blacklist},
        ip_address=request.META.get("REMOTE_ADDR"),
    )
    return redirect("clients_list")


@login_required
def add_user(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    if request.method == "POST":
        login = request.POST.get("login")
        password = request.POST.get("password")
        lname = request.POST.get("lname")
        fname = request.POST.get("fname")
        patronimyc = request.POST.get("patronimyc") or None
        phone = request.POST.get("phone")
        role = request.POST.get("role", "dispatcher")
        is_active = request.POST.get("is_active") == "on"

        user_obj = User.objects.create_user(
            login=login,
            password=password,
            lname=lname,
            fname=fname,
            patronimyc=patronimyc,
            phone=phone,
            role=role,
            is_active=is_active,
        )

        log_action(
            user=request.user,
            action="CREATE",
            object_type="User",
            object_id=user_obj.user_id,
            description=f"Создан пользователь {user_obj.lname} {user_obj.fname}",
            ip_address=request.META.get("REMOTE_ADDR"),
        )
        return redirect("users_list")

    return render(request, "add_user.html")


@login_required
def edit_user(request, user_id):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    user_obj = User.objects.get(pk=user_id)

    if request.method == "POST":
        old_data = {
            "login": user_obj.login,
            "lname": user_obj.lname,
            "fname": user_obj.fname,
            "patronimyc": user_obj.patronimyc,
            "phone": user_obj.phone,
            "role": user_obj.role,
            "is_active": user_obj.is_active,
        }
        user_obj.login = request.POST.get("login")
        user_obj.lname = request.POST.get("lname")
        user_obj.fname = request.POST.get("fname")
        user_obj.patronimyc = request.POST.get("patronimyc") or None
        user_obj.phone = request.POST.get("phone")
        user_obj.role = request.POST.get("role", user_obj.role)
        user_obj.is_active = request.POST.get("is_active") == "on"

        new_password = request.POST.get("password")
        if new_password:
            user_obj.set_password(new_password)

        user_obj.save()
        log_action(
            user=request.user,
            action="UPDATE",
            object_type="User",
            object_id=user_obj.user_id,
            description=f"Изменены данные пользователя {user_obj.lname} {user_obj.fname}",
            extra_data={"old_data": old_data},
            ip_address=request.META.get("REMOTE_ADDR"),
        )
        return redirect("users_list")

    return render(request, "edit_user.html", {"user_obj": user_obj})

@login_required
def toggle_user_archive(request, user_id):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    user_obj = User.objects.get(pk=user_id)
    old_archive = user_obj.is_archive
    user_obj.is_archive = not user_obj.is_archive
    user_obj.save()

    log_action(
        user=request.user,
        action="TOGGLE_ARCHIVE",
        object_type="User",
        object_id=user_obj.user_id,
        description=f"Архивирование пользователя {user_obj.lname} {user_obj.fname} ({'в архив' if user_obj.is_archive else 'из архива'})",
        extra_data={"old_archive": old_archive, "new_archive": user_obj.is_archive},
        ip_address=request.META.get("REMOTE_ADDR"),
    )

    return redirect("users_list")



@login_required
def add_shift(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    if request.method == "POST":
        driver_id = request.POST.get("driver")
        start_shift = request.POST.get("start_shift")  
        end_shift = request.POST.get("end_shift")


        shift = Shift.objects.create(
            driver_id=driver_id,
            start_shift=start_shift.replace("T", " "),
            end_shift=end_shift.replace("T", " "),
        )

        log_action(
            user=request.user,
            action="CREATE",
            object_type="Shift",
            object_id=shift.shift_id,
            description=f"Создана смена для водителя {shift.driver}",
            ip_address=request.META.get("REMOTE_ADDR"),
        )
        return redirect("shifts_list")

    drivers = Driver.objects.filter(is_archive=False)
    return render(request, "add_shift.html", {"drivers": drivers})


@login_required
def edit_shift(request, shift_id):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    shift = Shift.objects.get(pk=shift_id)

    if request.method == "POST":
        old_data = {
            "driver_id": shift.driver_id,
            "start_shift": str(shift.start_shift),
            "end_shift": str(shift.end_shift),
            "opened_user_id": shift.opened_user_id,
            "closed_user_id": shift.closed_user_id,
        }
        driver_id = request.POST.get("driver")
        start_shift = request.POST.get("start_shift")
        end_shift = request.POST.get("end_shift")
        opened_user_id = request.POST.get("opened_user") or None
        closed_user_id = request.POST.get("closed_user") or None

        shift.driver_id = driver_id
        shift.start_shift = start_shift.replace("T", " ")
        shift.end_shift = end_shift.replace("T", " ")
        shift.opened_user_id = opened_user_id if opened_user_id else None
        shift.closed_user_id = closed_user_id if closed_user_id else None
        shift.save()

        log_action(
            user=request.user,
            action="UPDATE",
            object_type="Shift",
            object_id=shift.shift_id,
            description=f"Изменены данные смены {shift.shift_id}",
            extra_data={"old_data": old_data},
            ip_address=request.META.get("REMOTE_ADDR"),
        )

        return redirect("shifts_list")

    drivers = Driver.objects.filter(is_archive=False)
    users = User.objects.filter(is_archive=False)
    return render(request, "edit_shift.html", {
        "shift": shift,
        "drivers": drivers,
        "users": users,
    })

@login_required
def districts_list(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    districts = District.objects.all()
    return render(request, "districts_list.html", {"districts": districts})


@login_required
def add_district(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    if request.method == "POST":
        
        name = request.POST.get("name")
        base_coefficient = request.POST.get("base_coefficient")

        district = District.objects.create(
            name=name,
            base_coefficient=base_coefficient,
        )

        log_action(
            user=request.user,
            action="CREATE",
            object_type="District",
            object_id=district.district_id,
            description=f"Создан район {district.name}",
            ip_address=request.META.get("REMOTE_ADDR"),
        )
        return redirect("districts_list")

    return render(request, "add_district.html")


@login_required
def edit_district(request, district_id):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    district = District.objects.get(pk=district_id)

    if request.method == "POST":
        old_data = {
            "name": district.name,
            "base_coefficient": str(district.base_coefficient),
        }
        district.name = request.POST.get("name")
        district.base_coefficient = request.POST.get("base_coefficient")
        district.save()
        log_action(
            user=request.user,
            action="UPDATE",
            object_type="District",
            object_id=district.district_id,
            description=f"Изменены данные района {district.name}",
            extra_data={"old_data": old_data},
            ip_address=request.META.get("REMOTE_ADDR"),
        )
        return redirect("districts_list")

    return render(request, "edit_district.html", {"district": district})


@login_required
def streets_list(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    streets = Street.objects.select_related("district").all()

    query = request.GET.get("q", "").strip()
    if query:
        streets = streets.filter(
            Q(name__icontains=query)
        )

    return render(request, "streets_list.html", {
        "streets": streets,
        "query": query,
    })


@login_required
def add_street(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    if request.method == "POST":
        name = request.POST.get("name")
        district_id = request.POST.get("district")

        street = Street.objects.create(
            name=name,
            district_id=district_id,
        )

        log_action(
            user=request.user,
            action="CREATE",
            object_type="Street",
            object_id=street.street_id,
            description=f"Создана улица {street.name}",
            ip_address=request.META.get("REMOTE_ADDR"),
        )
        return redirect("streets_list")

    districts = District.objects.all()
    return render(request, "add_street.html", {"districts": districts})


@login_required
def edit_street(request, street_id):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    street = Street.objects.get(pk=street_id)

    if request.method == "POST":
        old_data = {
            "name": street.name,
            "district_id": street.district_id,
        }
        street.name = request.POST.get("name")
        street.district_id = request.POST.get("district")
        street.save()

        log_action(
            user=request.user,
            action="UPDATE",
            object_type="Street",
            object_id=street.street_id,
            description=f"Изменены данные улицы {street.name}",
            extra_data={"old_data": old_data},
            ip_address=request.META.get("REMOTE_ADDR"),
        )
        return redirect("streets_list")

    districts = District.objects.all()
    return render(request, "edit_street.html", {
        "street": street,
        "districts": districts,
    })


def logout_view(request):
    if request.user.is_authenticated:
        log_action(
            user=request.user,
            action="LOGOUT",
            description=f"Выход из системы: {request.user.login}",
            ip_address=request.META.get("REMOTE_ADDR"),
        )
    
    logout(request)
    return redirect("login")


@login_required
def tariffs_list(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    tariffs = Tariff.objects.all()
    return render(request, "tariffs_list.html", {"tariffs": tariffs})


@login_required
def add_tariff(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    if request.method == "POST":
        name = request.POST.get("name")
        price = request.POST.get("price")
        time_from = request.POST.get("time_from")
        time_to = request.POST.get("time_to")

        days_of_week = 0
        for day_value in request.POST.getlist("days_of_week"):
            days_of_week |= int(day_value)

        tariff = Tariff.objects.create(
            name=name,
            price=price,
            time_from=time_from,
            time_to=time_to,
            days_of_week=days_of_week,
        )

        log_action(
            user=request.user,
            action="CREATE",
            object_type="Tariff",
            object_id=tariff.tariff_id,
            description=f"Создан тариф {tariff.name}",
            ip_address=request.META.get("REMOTE_ADDR"),
        )
        return redirect("tariffs_list")

    return render(request, "add_tariff.html", {
        "days_choices": Tariff.DAYS_OF_WEEK_CHOICES,
    })


@login_required
def edit_tariff(request, tariff_id):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    tariff = Tariff.objects.get(pk=tariff_id)

    selected_days = []
    for day_value, day_label in Tariff.DAYS_OF_WEEK_CHOICES:
        if tariff.days_of_week & day_value:
            selected_days.append(day_value)

    if request.method == "POST":
        old_data = {
            "name": tariff.name,
            "price": str(tariff.price),
            "time_from": str(tariff.time_from),
            "time_to": str(tariff.time_to),
            "days_of_week": tariff.days_of_week,
        }

        tariff.name = request.POST.get("name")
        tariff.price = request.POST.get("price")
        tariff.time_from = request.POST.get("time_from")
        tariff.time_to = request.POST.get("time_to")

        days_of_week = 0
        for day_value in request.POST.getlist("days_of_week"):
            days_of_week |= int(day_value)
        tariff.days_of_week = days_of_week

        tariff.save()
        log_action(
            user=request.user,
            action="UPDATE",
            object_type="Tariff",
            object_id=tariff.tariff_id,
            description=f"Изменены данные тарифа {tariff.name}",
            extra_data={"old_data": old_data},
            ip_address=request.META.get("REMOTE_ADDR"),
        )
        return redirect("tariffs_list")

    return render(request, "edit_tariff.html", {
        "tariff": tariff,
        "days_choices": Tariff.DAYS_OF_WEEK_CHOICES,
        "selected_days": selected_days,
    })

@login_required
def toggle_tariff_archive(request, tariff_id):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    tariff = Tariff.objects.get(pk=tariff_id)
    old_archive = tariff.is_archive
    tariff.is_archive = not tariff.is_archive
    tariff.save()

    log_action(
        user=request.user,
        action="TOGGLE_ARCHIVE",
        object_type="Tariff",
        object_id=tariff.tariff_id,
        description=f"Архивирование тарифа {tariff.name} ({'в архив' if tariff.is_archive else 'из архива'})",
        extra_data={"old_archive": old_archive, "new_archive": tariff.is_archive},
        ip_address=request.META.get("REMOTE_ADDR"),
    )

    return redirect("tariffs_list")


@login_required
def orders_list(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    query = request.GET.get("q", "")
    status_filter = request.GET.get("status", "")
    driver_filter = request.GET.get("driver", "")
    date_from = request.GET.get("date_from", "")
    date_to = request.GET.get("date_to", "")

    orders = Order.objects.select_related(
        "client", "driver", "tariff", "street_from", "street_to", "created_by"
    ).all()

    if query:
        orders = orders.filter(
            models.Q(client__lname__icontains=query) |
            models.Q(client__fname__icontains=query) |
            models.Q(client__phone__icontains=query) |
            models.Q(order_number__icontains=query)
        )

    if status_filter:
        orders = orders.filter(status=status_filter)

    if driver_filter:
        orders = orders.filter(driver_id=driver_filter)

    if date_from:
        orders = orders.filter(created_at__date__gte=date_from)
    if date_to:
        orders = orders.filter(created_at__date__lte=date_to)

    drivers = Driver.objects.filter(is_archive=False)

    context = {
        "orders": orders,
        "query": query,
        "current_status": status_filter,
        "current_driver": driver_filter,
        "date_from": date_from,
        "date_to": date_to,
        "status_choices": Order.STATUS_CHOICES,
        "drivers": drivers,
    }

    return render(request, "orders_list.html", context)

@login_required
def edit_order(request, order_id):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    order = Order.objects.select_related(
        "client", "driver", "tariff", "street_from", "street_to"
    ).get(pk=order_id)

    if request.method == "POST":
        action = request.POST.get("action")

        if action == "assign_driver":
            driver_id = request.POST.get("driver")
            if driver_id:
                old_driver = order.driver
                order.driver = Driver.objects.get(pk=driver_id)
                if order.status == "Новый":
                    order.status = "Назначен водитель"
                order.save()

                log_action(
                    user=request.user,
                    action="ASSIGN_DRIVER",
                    object_type="Order",
                    object_id=order.order_id,
                    description=f"Назначен водитель {order.driver} на заказ {order.order_number}",
                    extra_data={"old_driver": str(old_driver), "new_driver": str(order.driver)},
                    ip_address=request.META.get("REMOTE_ADDR"),
                )
            return redirect("orders_list")

        elif action == "complete":
            order.status = "Завершен"
            order.completed_at = timezone.now()
            order.save()

            log_action(
                user=request.user,
                action="COMPLETE_ORDER",
                object_type="Order",
                object_id=order.order_id,
                description=f"Заказ {order.order_number} завершён",
                ip_address=request.META.get("REMOTE_ADDR"),
            )
            return redirect("orders_list")

        elif action == "cancel":
            order.status = "Отменен"
            order.save()
            log_action(
                user=request.user,
                action="CANCEL_ORDER",
                object_type="Order",
                object_id=order.order_id,
                description=f"Заказ {order.order_number} отменён",
                ip_address=request.META.get("REMOTE_ADDR"),
            )
            return redirect("orders_list")

        old_data = {
            "client": str(order.client),
            "street_from": str(order.street_from),
            "house_from": order.house_from,
            "street_to": str(order.street_to),
            "house_to": order.house_to,
            "tariff": str(order.tariff),
            "price": str(order.price),
        }

        client_id = request.POST.get("client")
        street_from_id = request.POST.get("street_from")
        house_from = request.POST.get("house_from")
        street_to_id = request.POST.get("street_to")
        house_to = request.POST.get("house_to")
        tariff_id = request.POST.get("tariff")

        order.client = Client.objects.get(pk=client_id)
        order.street_from = Street.objects.select_related("district").get(pk=street_from_id)
        order.house_from = house_from
        order.street_to = Street.objects.select_related("district").get(pk=street_to_id)
        order.house_to = house_to
        order.tariff = Tariff.objects.get(pk=tariff_id)

        
        coef_from = order.street_from.district.base_coefficient
        coef_to = order.street_to.district.base_coefficient
        order.price = order.tariff.price * ((coef_from + coef_to) / 2)

        order.save()
        log_action(
            user=request.user,
            action="UPDATE",
            object_type="Order",
            object_id=order.order_id,
            description=f"Изменены данные заказа {order.order_number}",
            extra_data={"old_data": old_data, "new_data": str(order)},
            ip_address=request.META.get("REMOTE_ADDR"),
        )
        return redirect("orders_list")

    clients = Client.objects.all()
    streets = Street.objects.select_related("district").all()
    tariffs = Tariff.objects.filter(is_archive=False)
    drivers = Driver.objects.filter(is_archive=False)

    return render(request, "edit_order.html", {
        "order": order,
        "clients": clients,
        "streets": streets,
        "tariffs": tariffs,
        "drivers": drivers,
    })


@login_required
def audit_logs_list(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    logs = AuditLog.objects.select_related("user").all()

    user_filter = request.GET.get("user", "")
    action_filter = request.GET.get("action", "")
    date_from = request.GET.get("date_from", "")
    date_to = request.GET.get("date_to", "")

    if user_filter:
        logs = logs.filter(user_id=user_filter)
    if action_filter:
        logs = logs.filter(action=action_filter)
    if date_from:
        logs = logs.filter(created_at__date__gte=date_from)
    if date_to:
        logs = logs.filter(created_at__date__lte=date_to)

    users = User.objects.filter(is_archive=False)

    return render(request, "audit_logs_list.html", {
        "logs": logs,
        "users": users,
        "current_user": user_filter,
        "current_action": action_filter,
        "date_from": date_from,
        "date_to": date_to,
        "action_choices": AuditLog.ACTION_CHOICES,
    })

@login_required
def orders_report_page(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    drivers = Driver.objects.filter(is_archive=False)

    return render(request, "orders_report.html", {
        "drivers": drivers,
        "status_choices": Order.STATUS_CHOICES,
    })


@login_required
def export_orders_report(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    date_from = request.GET.get("date_from", "")
    date_to = request.GET.get("date_to", "")
    status_filter = request.GET.get("status", "")
    driver_filter = request.GET.get("driver", "")

    orders = Order.objects.select_related(
        "client", "driver", "tariff", "street_from", "street_to", "created_by"
    ).all()

    if date_from:
        orders = orders.filter(created_at__date__gte=date_from)
    if date_to:
        orders = orders.filter(created_at__date__lte=date_to)
    if status_filter:
        orders = orders.filter(status=status_filter)
    if driver_filter:
        orders = orders.filter(driver_id=driver_filter)

   
    wb = Workbook()
    ws = wb.active
    ws.title = "Заказы"
  
    headers = [
        "Номер заказа", "Дата создания", "Статус", "Клиент", "Телефон",
        "Водитель", "Авто", "Откуда", "Дом", "Куда", "Дом",
        "Тариф", "Цена", "Создал", "Время выполнения (мин)"
    ]
    ws.append(headers)

    
    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True)
        cell.fill = openpyxl.styles.PatternFill(start_color="4472C4", fill_type="solid")
        cell.font = openpyxl.styles.Font(color="FFFFFF", bold=True)

    
    for order in orders:
        duration = ""
        if order.completed_at and order.created_at:
            duration = int((order.completed_at - order.created_at).total_seconds() / 60)

        ws.append([
            order.order_number,
            order.created_at.strftime("%d.%m.%Y %H:%M"),
            order.status,
            f"{order.client.lname} {order.client.fname}",
            order.client.phone,
            f"{order.driver.lname} {order.driver.fname}" if order.driver else "",
            f"{order.driver.car_model} {order.driver.car_number}" if order.driver else "",
            order.street_from.name,
            order.house_from,
            order.street_to.name,
            order.house_to,
            order.tariff.name,
            str(order.price),
            f"{order.created_by.lname} {order.created_by.fname}",
            duration,
        ])

    
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = f"orders_report_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response["Content-Disposition"] = f"attachment; filename={filename}"
    wb.save(response)
    return response

@login_required
def dispatchers_report_page(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    return render(request, "dispatchers_report.html", {
        "roles": User.ROLE_CHOICES,
    })


@login_required
def export_dispatchers_report(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    
    role_filter = request.GET.get("role", "")
    is_active_filter = request.GET.get("is_active", "")
    is_archive_filter = request.GET.get("is_archive", "")
    date_from = request.GET.get("date_from", "")
    date_to = request.GET.get("date_to", "")

    users = User.objects.all()

    if role_filter:
        users = users.filter(role=role_filter)
    if is_active_filter == "1":
        users = users.filter(is_active=True)
    elif is_active_filter == "0":
        users = users.filter(is_active=False)
    if is_archive_filter == "1":
        users = users.filter(is_archive=True)
    elif is_archive_filter == "0":
        users = users.filter(is_archive=False)

    
    wb = Workbook()
    ws = wb.active
    ws.title = "Диспетчеры"

    headers = [
        "ID", "Логин", "Фамилия", "Имя", "Отчество", "Телефон",
        "Роль", "Активен", "Архив", "Создано заказов",
        "Открыто смен", "Закрыто смен", "Последний вход"
    ]
    ws.append(headers)


    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    for user in users:
        orders_count = Order.objects.filter(created_by=user).count()
        
        
        opened_shifts = Shift.objects.filter(opened_user=user).count()
        closed_shifts = Shift.objects.filter(closed_user=user).count()

        ws.append([
            user.user_id,
            user.login,
            user.lname,
            user.fname,
            user.patronimyc or "",
            user.phone,
            user.role,
            "Да" if user.is_active else "Нет",
            "Да" if user.is_archive else "Нет",
            orders_count,
            opened_shifts,
            closed_shifts,
            "",  
        ])

    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = f"dispatchers_report_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response["Content-Disposition"] = f"attachment; filename={filename}"
    wb.save(response)
    return response


@login_required
def drivers_report_page(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    return render(request, "drivers_report.html", {
        "statuses": Driver.STATUS_CHOICES,
    })


@login_required
def export_drivers_report(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    status_filter = request.GET.get("status", "")
    is_blacklist_filter = request.GET.get("is_blacklist", "")
    is_archive_filter = request.GET.get("is_archive", "")

    drivers = Driver.objects.all()

    if status_filter:
        drivers = drivers.filter(status=status_filter)
    if is_blacklist_filter == "1":
        drivers = drivers.filter(is_blacklist=True)
    elif is_blacklist_filter == "0":
        drivers = drivers.filter(is_blacklist=False)
    if is_archive_filter == "1":
        drivers = drivers.filter(is_archive=True)
    elif is_archive_filter == "0":
        drivers = drivers.filter(is_archive=False)

    wb = Workbook()
    ws = wb.active
    ws.title = "Водители"

    headers = [
        "ID", "Фамилия", "Имя", "Отчество", "Телефон",
        "Авто (модель)", "Авто (номер)", "Статус", "Чёрный список", "Архив",
        "Всего заказов", "Завершено заказов", "Выручка"
    ]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    for driver in drivers:
        total_orders = Order.objects.filter(driver=driver).count()
        completed_orders = Order.objects.filter(driver=driver, status="Завершен").count()
        
        revenue = Order.objects.filter(
            driver=driver, 
            status="Завершен"
        ).aggregate(total=models.Sum("price"))["total"] or 0

        ws.append([
            driver.driver_id,
            driver.lname,
            driver.fname,
            driver.patronimyc or "",
            driver.phone,
            driver.car_model,
            driver.car_number,
            driver.status,
            "Да" if driver.is_blacklist else "Нет",
            "Да" if driver.is_archive else "Нет",
            total_orders,
            completed_orders,
            str(revenue),
        ])

    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = f"drivers_report_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response["Content-Disposition"] = f"attachment; filename={filename}"
    wb.save(response)
    return response 

@login_required
def clients_report_page(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    return render(request, "clients_report.html", {})


@login_required
def export_clients_report(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    is_blacklist_filter = request.GET.get("is_blacklist", "")
    min_orders = request.GET.get("min_orders", "")

    clients = Client.objects.all()

    if is_blacklist_filter == "1":
        clients = clients.filter(is_blacklist=True)
    elif is_blacklist_filter == "0":
        clients = clients.filter(is_blacklist=False)

    if min_orders and min_orders.isdigit():
        min_orders = int(min_orders)
        clients = clients.annotate(
            orders_count=models.Count("orders")
        ).filter(orders_count__gte=min_orders)

    wb = Workbook()
    ws = wb.active
    ws.title = "Клиенты"

    headers = [
        "ID", "Фамилия", "Имя", "Отчество", "Телефон", "Чёрный список",
        "Всего заказов", "Последний заказ", "Общая сумма"
    ]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    for client in clients:
        total_orders = Order.objects.filter(client=client).count()
        
        last_order = Order.objects.filter(client=client).order_by("-created_at").first()
        last_order_date = last_order.created_at.strftime("%d.%m.%Y %H:%M") if last_order else ""
        
        total_amount = Order.objects.filter(
            client=client, 
            status="Завершен"
        ).aggregate(total=models.Sum("price"))["total"] or 0

        ws.append([
            client.client_id,
            client.lname,
            client.fname,
            client.patronimyc or "",
            client.phone,
            "Да" if client.is_blacklist else "Нет",
            total_orders,
            last_order_date,
            str(total_amount),
        ])

    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = f"clients_report_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response["Content-Disposition"] = f"attachment; filename={filename}"
    wb.save(response)
    return response

@login_required
def shifts_report_page(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    return render(request, "shifts_report.html", {
        "statuses": Shift.STATUS_CHOICES,
    })


@login_required
def export_shifts_report(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    status_filter = request.GET.get("status", "")
    driver_filter = request.GET.get("driver", "")

    shifts = Shift.objects.select_related("driver").all()

    if status_filter:
        shifts = shifts.filter(status=status_filter)
    if driver_filter:
        shifts = shifts.filter(driver__driver_id=driver_filter)

    wb = Workbook()
    ws = wb.active
    ws.title = "Смены"

    headers = [
        "ID", "Водитель (ФИО)", "Водитель (телефон)", "Авто (модель)", "Авто (номер)",
        "Начало", "Конец", "Длительность (часы)", "Статус",
        "Заказов", "Выручка"
    ]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    for shift in shifts:
        orders = Order.objects.filter(shift=shift)
        orders_count = orders.count()
        revenue = orders.aggregate(total=models.Sum("price"))["total"] or 0

        if shift.started_at and shift.ended_at:
            duration = (shift.ended_at - shift.started_at).total_seconds() / 3600
            duration_str = f"{duration:.1f}"
        else:
            duration_str = ""

        ws.append([
            shift.shift_id,
            f"{shift.driver.lname} {shift.driver.fname} {shift.driver.patronimyc or ''}",
            shift.driver.phone,
            shift.driver.car_model,
            shift.driver.car_number,
            shift.started_at.strftime("%d.%m.%Y %H:%M") if shift.started_at else "",
            shift.ended_at.strftime("%d.%m.%Y %H:%M") if shift.ended_at else "",
            duration_str,
            shift.status,
            orders_count,
            str(revenue),
        ])

    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = f"shifts_report_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response["Content-Disposition"] = f"attachment; filename={filename}"
    wb.save(response)
    return response

@login_required
def revenue_report_page(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    return render(request, "revenue_report.html", {})


@login_required
def export_revenue_report(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    period = request.GET.get("period", "day")  
    date_from = request.GET.get("date_from", "")
    date_to = request.GET.get("date_to", "")

    orders = Order.objects.filter(status="Завершен")

    if date_from:
        try:
            date_from_obj = datetime.strptime(date_from, "%Y-%m-%d")
            orders = orders.filter(created_at__gte=date_from_obj)
        except:
            pass
    if date_to:
        try:
            date_to_obj = datetime.strptime(date_to, "%Y-%m-%d")
            date_to_obj = date_to_obj + timedelta(days=1)
            orders = orders.filter(created_at__lt=date_to_obj)
        except:
            pass

    wb = Workbook()
    ws = wb.active
    ws.title = "Доходы"

    headers = [
        "Период", "Заказов", "Выручка", "Средний чек"
    ]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    if period == "day":
        from django.db.models import Func, F, CharField

        orders = orders.annotate(
            period=Func(
                F("created_at"),
                function="DATE",
                output_field=CharField(),
            )
        ).values("period").annotate(
            orders_count=Count("order_id"),
            revenue=Sum("price"),
        ).order_by("period")

        for row in orders:
            avg_check = row["revenue"] / row["orders_count"] if row["orders_count"] > 0 else 0
            ws.append([
                row["period"],
                row["orders_count"],
                str(row["revenue"] or 0),
                f"{avg_check:.2f}",
            ])

    elif period == "month":
        from django.db.models import Func, F, IntegerField

        orders = orders.annotate(
            year=Func(F("created_at"), function="YEAR", output_field=IntegerField()),
            month=Func(F("created_at"), function="MONTH", output_field=IntegerField()),
        ).values("year", "month").annotate(
            orders_count=Count("order_id"),
            revenue=Sum("price"),
        ).order_by("year", "month")

        for row in orders:
            avg_check = row["revenue"] / row["orders_count"] if row["orders_count"] > 0 else 0
            period_str = f"{row['month']:02d}.{row['year']}"
            ws.append([
                period_str,
                row["orders_count"],
                str(row["revenue"] or 0),
                f"{avg_check:.2f}",
            ])

    elif period == "year":
        from django.db.models import Func, F, IntegerField

        orders = orders.annotate(
            year=Func(F("created_at"), function="YEAR", output_field=IntegerField()),
        ).values("year").annotate(
            orders_count=Count("order_id"),
            revenue=Sum("price"),
        ).order_by("year")

        for row in orders:
            avg_check = row["revenue"] / row["orders_count"] if row["orders_count"] > 0 else 0
            ws.append([
                str(row["year"]),
                row["orders_count"],
                str(row["revenue"] or 0),
                f"{avg_check:.2f}",
            ])

    else:
        total_orders = orders.count()
        total_revenue = orders.aggregate(total=Sum("price"))["total"] or 0
        avg_check = total_revenue / total_orders if total_orders > 0 else 0

        ws.append([
            "Всего",
            total_orders,
            str(total_revenue),
            f"{avg_check:.2f}",
        ])

    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = f"revenue_report_{tz.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response["Content-Disposition"] = f"attachment; filename={filename}"
    wb.save(response)
    return response


@login_required
def audit_report_page(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    actions = AuditLog.ACTION_CHOICES
    users = User.objects.all()

    return render(request, "audit_report.html", {
        "actions": actions,
        "users": users,
    })


@login_required
def export_audit_report(request):
    if not hasattr(request.user, "role") or request.user.role != "admin":
        return HttpResponse("Доступ запрещён", status=403)

    action_filter = request.GET.get("action", "")
    user_filter = request.GET.get("user", "")
    object_type_filter = request.GET.get("object_type", "")
    date_from = request.GET.get("date_from", "")
    date_to = request.GET.get("date_to", "")

    audits = AuditLog.objects.select_related("user").all()

    if action_filter:
        audits = audits.filter(action=action_filter)
    if user_filter:
        audits = audits.filter(user__user_id=user_filter)
    if object_type_filter:
        audits = audits.filter(object_type__icontains=object_type_filter)
    if date_from:
        try:
            date_from_obj = datetime.strptime(date_from, "%Y-%m-%d")
            audits = audits.filter(created_at__gte=date_from_obj)
        except:
            pass
    if date_to:
        try:
            date_to_obj = datetime.strptime(date_to, "%Y-%m-%d")
            date_to_obj = date_to_obj + timedelta(days=1)
            audits = audits.filter(created_at__lt=date_to_obj)
        except:
            pass

    wb = Workbook()
    ws = wb.active
    ws.title = "Аудит"

    headers = [
        "ID", "Время", "Пользователь", "Действие", "Тип объекта", "ID объекта", "Описание", "Доп. данные", "IP"
    ]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    for audit in audits:
        user_name = ""
        if audit.user:
            user_name = f"{audit.user.lname} {audit.user.fname}"
        else:
            user_name = "Система"

        ws.append([
            audit.log_id,
            audit.created_at.strftime("%d.%m.%Y %H:%M:%S"),
            user_name,
            audit.get_action_display(),
            audit.object_type or "",
            audit.object_id or "",
            audit.description[:200] if audit.description else "",
            audit.extra_data[:200] if audit.extra_data else "",
            audit.ip_address or "",
        ])

    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = f"audit_report_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response["Content-Disposition"] = f"attachment; filename={filename}"
    wb.save(response)
    return response