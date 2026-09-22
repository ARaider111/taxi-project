import openpyxl
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.contrib.auth.decorators import login_required
from taxi.models import Driver, Client, Shift, Street, Tariff, Order
from django.utils import timezone
from django.db.models import Q
from openpyxl import Workbook
from django.db import models
from taxi.utils import log_action
from openpyxl.styles import Font, PatternFill, Alignment

@login_required
def dispatcher_dashboard(request):
    if not hasattr(request.user, "role") or request.user.role != "dispatcher":
        return HttpResponse("Доступ запрещён", status=403)
    return render(request, "dispatcher/dashboard.html")


@login_required
def dispatcher_drivers(request):
    if not hasattr(request.user, "role") or request.user.role != "dispatcher":
        return HttpResponse("Доступ запрещён", status=403)

    drivers = Driver.objects.all()

    query = request.GET.get("q", "").strip()
    if query:
        drivers = drivers.filter(
            Q(lname__icontains=query) |
            Q(fname__icontains=query) |
            Q(phone__icontains=query) |
            Q(car_model__icontains=query) |
            Q(car_number__icontains=query)
        )

    status = request.GET.get("status", "")
    if status:
        drivers = drivers.filter(status=status)

    statuses = Driver.STATUS_CHOICES  

    return render(request, "dispatcher/drivers_list.html", {
        "drivers": drivers,
        "statuses": statuses,
        "current_status": status,
        "query": query,
    })

@login_required
def dispatcher_clients(request):
    if not hasattr(request.user, "role") or request.user.role != "dispatcher":
        return HttpResponse("Доступ запрещён", status=403)

    clients = Client.objects.all()

    query = request.GET.get("q", "").strip()
    if query:
        clients = clients.filter(
            Q(lname__icontains=query) |
            Q(fname__icontains=query) |
            Q(phone__icontains=query)
        )

    return render(request, "dispatcher/clients_list.html", {
        "clients": clients,
        "query": query,
    })



def dispatcher_shifts(request):
    if not hasattr(request.user, "role") or request.user.role != "dispatcher":
        return HttpResponse("Доступ запрещён", status=403)

    shifts = Shift.objects.select_related(
        "driver", "opened_user", "closed_user"
    ).order_by("-start_shift")

    query = request.GET.get("q", "").strip()
    if query:
        shifts = shifts.filter(
            Q(driver__lname__icontains=query) |
            Q(driver__fname__icontains=query) |
            Q(start_shift__icontains=query) |
            Q(end_shift__icontains=query)
        )

  
    status = request.GET.get("status", "")
    if status == "not_opened":
        shifts = shifts.filter(opened_user__isnull=True)
    elif status == "opened":
        shifts = shifts.filter(
            opened_user__isnull=False,
            closed_user__isnull=True
        )
    elif status == "closed":
        shifts = shifts.filter(closed_user__isnull=False)

    return render(request, "dispatcher/shifts_list.html", {
        "shifts": shifts,
        "query": query,
        "current_status": status,
    })


@login_required
def dispatcher_add_client(request):
    if not hasattr(request.user, "role") or request.user.role != "dispatcher":
        return HttpResponse("Доступ запрещён", status=403)

    if request.method == "POST":
        lname = request.POST.get("lname")
        fname = request.POST.get("fname")
        patronimyc = request.POST.get("patronimyc") or None
        phone = request.POST.get("phone")

        client = Client.objects.create(
            lname=lname,
            fname=fname,
            patronimyc=patronimyc,
            phone=phone,
            is_blacklist=False,
        )

        log_action(
            user=request.user,
            action="CREATE",
            object_type="Client",
            object_id=client.client_id,
            description=f"Создан клиент {client.lname} {client.fname}",
            ip_address=request.META.get("REMOTE_ADDR"),
        )
        return redirect("dispatcher_clients")

    return render(request, "dispatcher/add_client.html")


@login_required
def dispatcher_edit_client(request, client_id):
    if not hasattr(request.user, "role") or request.user.role != "dispatcher":
        return HttpResponse("Доступ запрещён", status=403)

    client = Client.objects.get(pk=client_id)

    if request.method == "POST":
        old_data = {
            "lname": client.lname,
            "fname": client.fname,
            "patronimyc": client.patronimyc,
            "phone": client.phone,
        }

        client.lname = request.POST.get("lname")
        client.fname = request.POST.get("fname")
        client.patronimyc = request.POST.get("patronimyc") or None
        client.phone = request.POST.get("phone")
        client.save()

        log_action(
            user=request.user,
            action="UPDATE",
            object_type="Client",
            object_id=client.client_id,
            description=f"Изменены данные клиента {client.lname} {client.fname}",
            extra_data={"old_data": old_data},
            ip_address=request.META.get("REMOTE_ADDR"),
        )
        return redirect("dispatcher_clients")

    return render(request, "dispatcher/edit_client.html", {"client": client})

@login_required
def dispatcher_open_shift(request, shift_id):
    if not hasattr(request.user, "role") or request.user.role != "dispatcher":
        return HttpResponse("Доступ запрещён", status=403)

    shift = Shift.objects.get(pk=shift_id)

    if shift.opened_user is None:
        shift.start_shift = timezone.now()
        shift.opened_user = request.user
        shift.save()

        log_action(
            user=request.user,
            action="UPDATE",
            object_type="Shift",
            object_id=shift.shift_id,
            description=f"Открыта смена {shift.shift_id} для водителя {shift.driver}",
            ip_address=request.META.get("REMOTE_ADDR"),
        )

    return redirect("dispatcher_shifts")


@login_required
def dispatcher_close_shift(request, shift_id):
    if not hasattr(request.user, "role") or request.user.role != "dispatcher":
        return HttpResponse("Доступ запрещён", status=403)

    shift = Shift.objects.get(pk=shift_id)

    if shift.opened_user is not None and shift.closed_user is None:
        shift.end_shift = timezone.now()
        shift.closed_user = request.user
        shift.save()

        log_action(
            user=request.user,
            action="UPDATE",
            object_type="Shift",
            object_id=shift.shift_id,
            description=f"Закрыта смена {shift.shift_id} для водителя {shift.driver}",
            ip_address=request.META.get("REMOTE_ADDR"),
        )

    return redirect("dispatcher_shifts")



@login_required
def dispatcher_orders_list(request):
    if not hasattr(request.user, "role") or request.user.role not in ["admin", "dispatcher"]:
        return HttpResponse("Доступ запрещён", status=403)

    query = request.GET.get("q", "")
    status_filter = request.GET.get("status", "")

    orders = Order.objects.select_related(
        "client", "driver", "tariff", "street_from", "street_to", "created_by"
    ).all()

    if query:
        orders = orders.filter(
            models.Q(client__lname__icontains=query) |
            models.Q(client__fname__icontains=query) |
            models.Q(client__phone__icontains=query) |
            models.Q(order_number__icontains=query)
        )


    if status_filter:
        orders = orders.filter(status=status_filter)

    context = {
        "orders": orders,
        "query": query,
        "current_status": status_filter,
        "status_choices": Order.STATUS_CHOICES,
    }

    return render(request, "dispatcher/orders_list.html", context)


@login_required
def dispatcher_add_order(request):
    if not hasattr(request.user, "role") or request.user.role not in ["admin", "dispatcher"]:
        return HttpResponse("Доступ запрещён", status=403)

    if request.method == "POST":
        client_id = request.POST.get("client")
        street_from_id = request.POST.get("street_from")
        house_from = request.POST.get("house_from")
        street_to_id = request.POST.get("street_to")
        house_to = request.POST.get("house_to")
        tariff_id = request.POST.get("tariff")

        client = Client.objects.get(pk=client_id)
        street_from = Street.objects.select_related("district").get(pk=street_from_id)
        street_to = Street.objects.select_related("district").get(pk=street_to_id)
        tariff = Tariff.objects.get(pk=tariff_id)

        coef_from = street_from.district.base_coefficient
        coef_to = street_to.district.base_coefficient
        price = tariff.price * ((coef_from + coef_to) / 2)

        order = Order.objects.create(
            client=client,
            street_from=street_from,
            house_from=house_from,
            street_to=street_to,
            house_to=house_to,
            tariff=tariff,
            price=price,
            created_by=request.user,
            status="Новый",
        )

        log_action(
            user=request.user,
            action="CREATE",
            object_type="Order",
            object_id=order.order_id,
            description=f"Создан заказ {order.order_number} для клиента {client}",
            ip_address=request.META.get("REMOTE_ADDR"),
        )

        return redirect("dispatcher_orders_list")

   
    clients = Client.objects.all()
    streets = Street.objects.select_related("district").all()
    tariffs = Tariff.objects.filter(is_archive=False)

    return render(request, "dispatcher/add_order.html", {
        "clients": clients,
        "streets": streets,
        "tariffs": tariffs,
    })


@login_required
def dispatcher_edit_order(request, order_id):
    if not hasattr(request.user, "role") or request.user.role not in ["admin", "dispatcher"]:
        return HttpResponse("Доступ запрещён", status=403)

    order = Order.objects.select_related(
        "client", "driver", "tariff", "street_from", "street_to"
    ).get(pk=order_id)

    if order.status == "Завершен":
        return HttpResponse("Нельзя редактировать завершённый заказ", status=403)

    if request.method == "POST":
        old_data = {
            "client": str(order.client),
            "street_from": str(order.street_from),
            "house_from": order.house_from,
            "street_to": str(order.street_to),
            "house_to": order.house_to,
            "tariff": str(order.tariff),
            "price": str(order.price),
        }
         
        client_id = request.POST.get("client")
        street_from_id = request.POST.get("street_from")
        house_from = request.POST.get("house_from")
        street_to_id = request.POST.get("street_to")
        house_to = request.POST.get("house_to")
        tariff_id = request.POST.get("tariff")

        order.client = Client.objects.get(pk=client_id)
        order.street_from = Street.objects.select_related("district").get(pk=street_from_id)
        order.house_from = house_from
        order.street_to = Street.objects.select_related("district").get(pk=street_to_id)
        order.house_to = house_to
        order.tariff = Tariff.objects.get(pk=tariff_id)

        coef_from = order.street_from.district.base_coefficient
        coef_to = order.street_to.district.base_coefficient
        order.price = order.tariff.price * ((coef_from + coef_to) / 2)

        order.save()
        log_action(
            user=request.user,
            action="UPDATE",
            object_type="Order",
            object_id=order.order_id,
            description=f"Изменены данные заказа {order.order_number}",
            extra_data={"old_data": old_data},
            ip_address=request.META.get("REMOTE_ADDR"),
        )

        return redirect("orders_list")

    clients = Client.objects.all()
    streets = Street.objects.select_related("district").all()
    tariffs = Tariff.objects.filter(is_archive=False)

    return render(request, "dispatcher/edit_order.html", {
        "order": order,
        "clients": clients,
        "streets": streets,
        "tariffs": tariffs,
    })


@login_required
def dispatcher_assign_driver(request, order_id):
    if not hasattr(request.user, "role") or request.user.role not in ["admin", "dispatcher"]:
        return HttpResponse("Доступ запрещён", status=403)

    order = Order.objects.get(pk=order_id)

    if order.status == "Завершен":
        return HttpResponse("Нельзя изменить завершённый заказ", status=403)

    if request.method == "POST":
        driver_id = request.POST.get("driver")
        driver = Driver.objects.get(pk=driver_id)

        old_driver = order.driver
        order.driver = driver
        order.status = "Назначен водитель"
        order.save()
        log_action(
            user=request.user,
            action="ASSIGN_DRIVER",
            object_type="Order",
            object_id=order.order_id,
            description=f"Назначен водитель {order.driver} на заказ {order.order_number}",
            extra_data={"old_driver": str(old_driver), "new_driver": str(order.driver)},
            ip_address=request.META.get("REMOTE_ADDR"),
        )

        return redirect("orders_list")

    drivers = Driver.objects.filter(status="В ожидании", is_blacklist=False, is_archive=False)

    return render(request, "dispatcher/assign_driver.html", {
        "order": order,
        "drivers": drivers,
    })



@login_required
def dispatcher_complete_order(request, order_id):
    if not hasattr(request.user, "role") or request.user.role not in ["admin", "dispatcher"]:
        return HttpResponse("Доступ запрещён", status=403)

    order = Order.objects.get(pk=order_id)

    if order.status == "Завершен":
        return HttpResponse("Заказ уже завершён", status=403)

    order.status = "Завершен"
    order.completed_at = timezone.now()
    order.save()

    log_action(
        user=request.user,
        action="COMPLETE_ORDER",
        object_type="Order",
        object_id=order.order_id,
        description=f"Заказ {order.order_number} завершён",
        ip_address=request.META.get("REMOTE_ADDR"),
    )

    return redirect("orders_list")


@login_required
def dispatcher_shift_report(request):
    if not hasattr(request.user, "role") or request.user.role not in ["admin", "dispatcher"]:
        return HttpResponse("Доступ запрещён", status=403)

    user_shifts = Shift.objects.filter(
        models.Q(opened_user=request.user) | models.Q(closed_user=request.user)
    ).distinct()

    return render(request, "dispatcher/dispatcher_shift_report.html", {
        "user_shifts": user_shifts,
    })


@login_required
def export_dispatcher_shift_report(request):

    if not hasattr(request.user, "role") or request.user.role not in ["admin", "dispatcher"]:
        return HttpResponse("Доступ запрещён", status=403)
    
    shift_filter = request.GET.get("shift", "")
   
    user_shifts = Shift.objects.filter(
        models.Q(opened_user=request.user) | models.Q(closed_user=request.user)
    ).distinct()

    if shift_filter:
        user_shifts = user_shifts.filter(shift_id=shift_filter)

    wb = Workbook()
    ws = wb.active
    ws.title = "Отчет диспетчера"

    headers = [
        "ID заказа", "Номер заказа", "Клиент", "Адрес откуда", "Адрес куда",
        "Водитель", "Тариф", "Цена", "Статус", "Время создания", "Время завершения"
    ]
    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    orders = Order.objects.filter(created_by=request.user).select_related(
        "client", "driver", "tariff", "street_from", "street_to"
    )

  
    if shift_filter and user_shifts.exists():
        shift = user_shifts.filter(shift_id=shift_filter).first()
        if shift:
            orders = orders.filter(
                created_at__gte=shift.start_shift,
                created_at__lte=shift.end_shift,
            )

    for order in orders:
        ws.append([
            order.order_id,
            order.order_number,
            f"{order.client.lname} {order.client.fname}",
            f"{order.street_from.name}, {order.house_from}",
            f"{order.street_to.name}, {order.house_to}",
            f"{order.driver.lname} {order.driver.fname}" if order.driver else "Не назначен",
            order.tariff.name,
            str(order.price),
            order.status,
            order.created_at.strftime("%d.%m.%Y %H:%M"),
            order.completed_at.strftime("%d.%m.%Y %H:%M") if order.completed_at else "",
        ])

    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = f"dispatcher_report_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response["Content-Disposition"] = f"attachment; filename={filename}"
    wb.save(response)
    return response